#!/usr/bin/env python3
"""Tests for "Build a tracker" (engine/nl_tracker.py through nl_browser.tracker_json, CONTRACT-v2 section 5.20).

    ../agent-demo/venv/bin/python tools/test_nl_tracker.py

The tracker's rows are the "Clean my file" view's rows (one cleaner), so these tests hold it to:
  * pandas, computed here independently from the two SYNTHETIC demo files (tools/make_sample_tracker.py): which rows
    are placed, merged within a file, merged across files, unplaced; every row of every file accounted for once
  * its promises: nothing invented (every Data row is a line of a file), nothing filled (a cell is filled only when
    the file's cell was), every changed cell listed with its value as written and its value in the tracker, a
    withheld column in no sheet
  * the workbook, written by the page's own writer (src/js/56-xlsx.js under node) and opened with openpyxl: the
    formulas are there, the dropdowns reference the Lists ranges that hold the list, the Data sheet is an Excel table
    whose formula columns are calculated columns; every formula's cached value equals this test's own computation
    from pandas, and (with LibreOffice) equals what LibreOffice computes when it recalculates the file
  * several files at the engine level (two runs in one process, both orders), one file alone, the kind suggested for
    other shapes (inventory, shifts, vendors), the visitor's questions, a duplicate key, the cleaning edits
"""
from __future__ import annotations

import csv
import datetime as dt
import glob
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile

os.environ.setdefault("NL_BROWSER_STRICT", "1")
HERE = os.path.dirname(os.path.abspath(__file__))
SITE = os.path.normpath(os.path.join(HERE, ".."))
ENGINE_ROOT = os.path.abspath(os.environ.get("NL_ENGINE_ROOT") or os.path.join(SITE, "..", "northledger-core"))
sys.path.insert(0, ENGINE_ROOT)
sys.path.insert(0, os.path.join(SITE, "engine"))
import pandas as pd  # noqa: E402
import nl_browser as NB  # noqa: E402
import nl_tracker as TR  # noqa: E402

A = os.path.join(SITE, "engine", "sample-tracker-tasks-a-synthetic.csv")
B = os.path.join(SITE, "engine", "sample-tracker-tasks-b-synthetic.csv")
AS_OF = "2026-04-15"
NA, NBN = os.path.basename(A), os.path.basename(B)
_RUNS: dict = {}


def _run(path: str, data: bytes = None, name: str = None, decisions=None) -> str:
    """Run a file through the engine (as the page does); returns its sha256. The tab's memory holds every run."""
    data = data if data is not None else open(path, "rb").read()
    name = name or os.path.basename(path)
    key = (data, name, json.dumps(decisions, sort_keys=True))
    if key not in _RUNS or _RUNS[key] not in NB._CLEANS:
        rep = NB.run(data, name, "", decisions, AS_OF)
        assert rep["ok"], rep.get("error")
        _RUNS[key] = rep["input"]["sha256"]
    return _RUNS[key]


def _tracker(shas, **opt) -> dict:
    o = dict(opt, sha256=shas[-1], inputs=list(shas))
    o.setdefault("today", AS_OF)
    out = json.loads(NB.tracker_json(json.dumps(o)))
    assert out["ok"], out.get("error")
    return out


def _pair(**opt) -> dict:
    a, b = _run(A), _run(B)
    return _tracker([a, b], **opt)


def _sheet(out, name):
    return next(s for s in out["workbook"]["sheets"] if s["name"] == name)


def _v(c):
    return c.get("v") if isinstance(c, dict) else c


def _data(out):
    rows = _sheet(out, "Data")["rows"]
    heads = [_v(c) for c in rows[0]]
    return heads, [dict(zip(heads, r)) for r in rows[1:]]


def _raw(path) -> pd.DataFrame:
    return pd.read_csv(path, dtype=str, keep_default_na=False)


# the test's own reading of the demo files (independent of the engine): its status words, ISO dates, trimmed text
STATUS = {"done": "Done", "complete": "Done", "closed": "Done", "resolved": "Done", "open": "Open", "new": "Open",
          "in progress": "In progress", "wip": "In progress", "on hold": "Waiting", "waiting": "Waiting",
          "cancelled": "Cancelled"}
COLS_A = {"Task ID": "ID", "Task": "Title", "Status": "Status", "Type": "Category", "Team": "Owner",
          "Date logged": "Opened", "Due": "Due", "Completed": "Closed", "Est hours": "Est hours"}
COLS_B = {"ID": "ID", "Title": "Title", "State": "Status", "Category": "Category", "Team": "Owner",
          "Opened": "Opened", "Due date": "Due", "Closed": "Closed"}          # Notes: withheld (personal-data step)


def _iso(v: str, dayfirst: bool) -> str:
    v = v.strip()
    if not v or v.upper() == "N/A":
        return ""
    try:
        return pd.to_datetime(v.replace("/", "-") if v[:4].isdigit() else v, dayfirst=dayfirst,
                              format="%Y-%m-%d" if v[:4].isdigit() else "%d/%m/%Y").strftime("%Y-%m-%d")
    except Exception:  # noqa: BLE001 - not a date: as written
        return v


def _expected():
    """Each file's rows as the test reads them: (file, line, {tracker column: value}) with not-a-record rows aside."""
    out, unplaced, merged_within = [], [], []
    for path, cols, dayfirst in ((A, COLS_A, True), (B, COLS_B, False)):
        df = _raw(path)
        seen = {}
        for i, r in df.iterrows():
            line = i + 2
            vals = [r[c] for c in df.columns]
            filled = [v for v in vals if v.strip()]
            if all(str(v).strip().lower() == str(h).strip().lower() for v, h in zip(vals, df.columns) if v.strip()) and len(filled) > 1:
                unplaced.append((os.path.basename(path), line))
                continue
            if len(filled) == 1:
                unplaced.append((os.path.basename(path), line))
                continue
            key = tuple(vals)
            if key in seen:
                merged_within.append((os.path.basename(path), line, seen[key]))
                continue
            seen[key] = line
            rec = {}
            for c, t in cols.items():
                v = r[c].strip()
                if t in ("Opened", "Due", "Closed"):
                    v = _iso(v, dayfirst)
                elif t == "Status":
                    v = STATUS.get(v.lower(), v)
                elif t == "Category":
                    v = v.upper()
                elif t == "Est hours":
                    v = v.replace(" hrs", "")
                rec[t] = v
            out.append((os.path.basename(path), line, rec))
    return out, unplaced, merged_within


# ------------------------------------------------------------------ the demo pair, against pandas
def test_rows_against_pandas():
    out = _pair()
    exp, unplaced, within = _expected()
    heads, data = _data(out)
    # the template's columns, the formula columns, the column no template names (Est hours), then where it came from
    assert heads[:8] == ["ID", "Title", "Status", "Category", "Owner", "Opened", "Due", "Closed"], heads
    assert heads[8:12] == ["Days open", "Is open", "Overdue", "Oldest-open sort key"], heads
    assert heads[12:] == ["Est hours", "Source file", "Line in source", "Check"], heads
    # across files: the same task in both files, equal in every column both hold once cleaned -> merged into A's copy
    a_by_id = {r["ID"]: r for f, _l, r in exp if f == NA}
    shared = [c for c in COLS_B.values()]
    across = [(l, r["ID"]) for f, l, r in exp if f == NBN and r["ID"] in a_by_id and
              all(a_by_id[r["ID"]][c] == r[c] for c in shared)]
    assert sorted(i for _l, i in across) == ["T-115", "T-116", "T-118", "T-122", "T-123", "T-124"], across
    placed = [(f, l) for f, l, r in exp if not (f == NBN and (l, r["ID"]) in across)]
    assert len(data) == len(placed) == 37, (len(data), len(placed))
    assert [(d["Source file"].split(" (also")[0], d["Line in source"]) for d in data] == placed
    S = out["summary"]
    assert S["merged_across"] == 6 and S["merged_within"] == 2 and S["unplaced"] == 2 and S["rows"] == 37, S
    un = _sheet(out, "Unplaced")["rows"][1:]
    assert sorted((r[0], r[1]) for r in un) == sorted(unplaced), un
    mg = _sheet(out, "Merged duplicates")["rows"][1:]
    assert sorted((r[3]) for r in mg if r[0].startswith("within")) == sorted("%s line %d" % (f, l) for f, l, _k in within)
    assert sorted(r[3] for r in mg if r[0].startswith("across")) == sorted("%s line %d" % (NBN, l) for l, _i in across)
    # every line of every file accounted for exactly once: placed, merged or unplaced
    for path in (A, B):
        n = len(_raw(path))
        name = os.path.basename(path)
        lines = sorted([l for f, l in placed if f == name] + [int(r[3].rsplit(" ", 1)[1]) for r in mg if r[3].startswith(name)] +
                       [l for f, l in unplaced if f == name])
        assert lines == list(range(2, n + 2)), (name, lines)
    # the values: the test's own reading equals the tracker's, cell for cell (a doubtful cell as written)
    by = {(f, l): r for f, l, r in exp}
    for d in data:
        e = by[(d["Source file"].split(" (also")[0], d["Line in source"])]
        for c, v in e.items():
            got = _v(d[c])
            got = "" if got is None else (("%g" % got) if isinstance(got, float) else str(got))
            assert got == v, (d["ID"], c, got, v)
    # T-117, T-119 and T-125 changed between the exports: both copies kept, each flagged with the other
    for i in ("T-117", "T-119", "T-125"):
        two = [d for d in data if _v(d["ID"]) == i]
        assert len(two) == 2 and all("possible duplicate: the same ID" in (d["Check"] or "") for d in two), two


def test_nothing_invented_or_filled():
    out = _pair()
    heads, data = _data(out)
    raw = {NA: _raw(A), NBN: _raw(B)}
    cols = {NA: {v: k for k, v in COLS_A.items()}, NBN: {v: k for k, v in COLS_B.items()}}
    blanks = 0
    for d in data:
        f, line = d["Source file"].split(" (also")[0], d["Line in source"]
        src = raw[f].iloc[line - 2]
        for c in heads[:8] + ["Est hours"]:
            h = cols[f].get(c)
            v = _v(d[c])
            if h is None:
                assert v in (None, ""), (f, line, c, v)        # a column this file does not have stays empty
                continue
            if v in (None, ""):
                blanks += 1
            else:
                assert src[h].strip() and src[h].strip().upper() != "N/A", (f, line, c, v, src[h])   # never filled
    # T-126: "done" with no completed date stays blank and is flagged
    t126 = next(d for d in data if _v(d["ID"]) == "T-126")
    assert t126["Closed"] in (None, "") and "Status is Done but Closed is blank; left blank" in t126["Check"]
    flags = _sheet(out, "Flags")["rows"][1:]
    assert any(r[0] == NA and r[1] == 29 and r[3] == "Closed" and "left blank" in r[5] for r in flags)
    # every blank cell of a column the row's file holds is flagged (Closed only when the status is not open)
    n_blank_flags = sum(len(str(r[3]).split("; ")) for r in flags if r[5] == "blank, left empty: nothing was filled in")
    assert 0 < n_blank_flags <= blanks
    # the doubtful due date stays as written, flagged with the cleaner's words
    t114 = next(d for d in data if _v(d["ID"]) == "T-114")
    assert _v(t114["Due"]) == "next tue" and "matches no known date format" in t114["Check"]


def test_changes_list_every_changed_cell():
    out = _pair()
    heads, data = _data(out)
    raw = {NA: _raw(A), NBN: _raw(B)}
    cols = {NA: {v: k for k, v in COLS_A.items()}, NBN: {v: k for k, v in COLS_B.items()}}
    ch = {(r[0], r[1], r[3]): r for r in _sheet(out, "Changes")["rows"][1:]}
    n = 0
    for d in data:
        f, line = d["Source file"].split(" (also")[0], d["Line in source"]
        src = raw[f].iloc[line - 2]
        for c in heads[:8] + ["Est hours"]:
            h = cols[f].get(c)
            if h is None:
                continue
            before = src[h]
            v = _v(d[c])
            after = "" if v is None else (("%g" % v) if isinstance(v, float) else str(v))
            if c in ("Opened", "Due", "Closed") and isinstance(d[c], dict):
                after = d[c]["v"]
            changed = after != before and not (after == "" and before.strip() == "")
            row = ch.get((f, line, c))
            if changed:
                n += 1
                assert row is not None, (f, line, c, before, after)
                assert row[4] == before and row[5] == after and row[2] == data.index(d) + 2, (row, before, after)
                assert row[6], row
            else:
                assert row is None, row
    assert n == len(ch) == out["summary"]["changes"], (n, len(ch))
    rules = {r[6] for r in ch.values()}
    assert "status_case, then tracker_status_words" in rules and "due_date" in rules and "est_hours_numeric" in rules, rules


def test_withheld_column_is_in_no_sheet():
    out = _pair()
    text = json.dumps(out["workbook"])
    assert "Notes" not in text and "Signed renewal" not in text and "Awaiting quote" not in text
    ins = [_v(r[0]) for r in _sheet(out, "Instructions")["rows"] if r]
    assert any("1 withheld column left out of %s" % NBN in str(x) for x in ins), ins


# ------------------------------------------------------------------ the workbook, written by the page's writer
NODE_WRITE = r"""
import fs from 'node:fs'; import vm from 'node:vm';
const ctx = { window: {}, TextEncoder, Date, Math, Object, Array, String, Number, isFinite, Uint8Array, JSON, RegExp };
vm.createContext(ctx);
vm.runInContext(fs.readFileSync(process.argv[2], 'utf8'), ctx);
fs.writeFileSync(process.argv[4], ctx.window.NLX.workbook(JSON.parse(fs.readFileSync(process.argv[3], 'utf8'))));
"""


def _xlsx(out, tmp) -> str:
    spec = os.path.join(tmp, "spec.json")
    js = os.path.join(tmp, "write.mjs")
    dst = os.path.join(tmp, "tracker.xlsx")
    with open(spec, "w") as f:
        json.dump(out["workbook"], f)
    with open(js, "w") as f:
        f.write(NODE_WRITE)
    r = subprocess.run(["node", js, os.path.join(SITE, "src", "js", "56-xlsx.js"), spec, dst], capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    return dst


def test_workbook_opens_with_formulas_dropdowns_table():
    import openpyxl
    out = _pair(questions="How many tasks are overdue?\nWhich open tasks are oldest?\nWhat does the weather do?")
    tmp = tempfile.mkdtemp(prefix="nltrk_")
    try:
        wb = openpyxl.load_workbook(_xlsx(out, tmp))
        assert wb.sheetnames == ["Data", "Summary", "Lists", "Instructions", "Changes", "Merged duplicates", "Unplaced", "Flags"]
        ws = wb["Data"]
        heads = [c.value for c in ws[1]]
        L = {h: openpyxl.utils.get_column_letter(i + 1) for i, h in enumerate(heads)}
        assert list(ws.tables) == ["Tracker"] and ws.tables["Tracker"].ref == "A1:%s38" % L["Check"]
        calc = {c.name: (c.calculatedColumnFormula.attr_text if c.calculatedColumnFormula is not None else None)
                for c in ws.tables["Tracker"].tableColumns}
        for h in ("Days open", "Is open", "Overdue", "Oldest-open sort key"):
            assert calc[h] and ws["%s2" % L[h]].value == "=" + calc[h], (h, calc[h], ws["%s2" % L[h]].value)
        assert calc["Days open"].startswith("IF(ISNUMBER(Tracker[[#This Row],[Opened]])")
        # dropdowns: Status and Category, each from the Lists range that holds exactly its list
        lists = wb["Lists"]
        dv = {str(d.sqref).split(":")[0].rstrip("0123456789"): d.formula1 for d in ws.data_validations.dataValidation}
        assert set(dv) == {L["Status"], L["Category"]}, dv
        lh = [c.value for c in lists[1]]
        for col, name in ((L["Status"], "Status"), (L["Category"], "Category")):
            j = lh.index(name)
            ref = dv[col].replace("'", "")
            letter = openpyxl.utils.get_column_letter(j + 1)
            vals = [lists.cell(r, j + 1).value for r in range(2, lists.max_row + 1) if lists.cell(r, j + 1).value]
            assert ref == "Lists!$%s$2:$%s$%d" % (letter, letter, len(vals) + 1), (ref, vals)
            used = {ws["%s%d" % (col, r)].value for r in range(2, 39)} - {None}
            assert used <= set(vals), (used, vals)
        assert ws["%s2" % L["Opened"]].number_format == "yyyy-mm-dd" and isinstance(ws["%s2" % L["Opened"]].value, dt.datetime)
        sm = wb["Summary"]
        f = {sm.cell(r, 1).value: sm.cell(r, 2).value for r in range(1, sm.max_row + 1)}
        assert f["Overdue (open, due before today)"] == '=COUNTIFS(Data!$%s:$%s,"Yes")' % (L["Overdue"], L["Overdue"])
        assert f["How many tasks are overdue?"].startswith("=B"), f["How many tasks are overdue?"]
        assert str(f["What does the weather do?"]).startswith("Not answered by a formula"), f
        assert str(f["Which open tasks are oldest?"]).startswith("=C"), f["Which open tasks are oldest?"]
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _expected_formulas(data, today: dt.date):
    """This test's own computation of the formula columns and the Summary's totals, from the Data values."""
    t = pd.Timestamp(today)
    rows = []
    for d in data:
        o = pd.Timestamp(d["Opened"]["v"]) if isinstance(d["Opened"], dict) else None
        c = pd.Timestamp(d["Closed"]["v"]) if isinstance(d["Closed"], dict) else None
        du = pd.Timestamp(d["Due"]["v"]) if isinstance(d["Due"], dict) else None
        st = _v(d["Status"]) or ""
        days = "" if o is None else float(((c if c is not None else t) - o).days) if (c is not None or _v(d["Closed"]) in (None, "")) else ""
        is_open = "" if not st else ("No" if st in ("Done", "Cancelled") else "Yes")
        over = ("Yes" if du < t else "No") if is_open == "Yes" and du is not None else ""
        rows.append({"days": days, "open": is_open, "over": over})
    return rows


def test_formula_values_against_pandas():
    out = _pair()
    heads, data = _data(out)
    exp = _expected_formulas(data, dt.date(2026, 4, 15))
    for d, e in zip(data, exp):
        assert _v(d["Days open"]) == e["days"], (d["ID"], _v(d["Days open"]), e)
        assert _v(d["Is open"]) == e["open"] and _v(d["Overdue"]) == e["over"], (d["ID"], e)
    sm = {(_v(r[0]) if r else None): r for r in _sheet(out, "Summary")["rows"] if r}
    assert _v(sm["Open (status not in the closed list)"][1]) == sum(e["open"] == "Yes" for e in exp) == 20
    assert _v(sm["Overdue (open, due before today)"][1]) == sum(e["over"] == "Yes" for e in exp) == 16
    closed = [e["days"] for e in exp if e["open"] == "No" and e["days"] != ""]
    assert abs(_v(sm["Average days from opened to closed (closed items)"][1]) - sum(closed) / len(closed)) < 1e-9
    # the oldest open task: T-103, opened 2026-01-08, 97 days to 2026-04-15
    oldest = [r for r in _sheet(out, "Summary")["rows"] if r and _v(r[0]) == 1][0]
    assert _v(oldest[1]) == 97 and _v(oldest[2]) == "Update tenant handbook" and _v(oldest[3]) == "T-103", oldest
    by_cat = {}
    for d, e in zip(data, exp):
        k = _v(d["Category"]) or "(blank)"
        by_cat.setdefault(k, [0, 0])
        by_cat[k][0] += 1
        by_cat[k][1] += e["open"] == "Yes"
    rows = _sheet(out, "Summary")["rows"]
    at = next(i for i, r in enumerate(rows) if r and _v(r[0]) == "By category")
    for r in rows[at + 1:at + 1 + len(by_cat)]:
        assert [_v(r[1]), _v(r[2])] == by_cat[_v(r[0])], (r, by_cat)


def test_libreoffice_recalculates_the_same():
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("  (no LibreOffice: the recalculation check was not run)")
        return
    today = dt.date.today()
    out = _pair(today=today.isoformat())
    tmp = tempfile.mkdtemp(prefix="nltrk_lo_")
    try:
        x = _xlsx(out, tmp)
        r = subprocess.run([soffice, "--headless", "-env:UserInstallation=file://%s/profile" % tmp, "--convert-to",
                            "csv:Text - txt - csv (StarCalc):44,34,76,1,,0,false,true,false,false,false,-1",
                            "--outdir", tmp, x], capture_output=True, text=True, timeout=240)
        files = {os.path.basename(p).split("-", 1)[1][:-4]: p for p in glob.glob(os.path.join(tmp, "tracker-*.csv"))}
        assert "Data" in files and "Summary" in files, (r.stdout, r.stderr, files)
        lo = list(csv.reader(open(files["Data"], encoding="utf-8")))
        heads, data = _data(out)
        exp = _expected_formulas(data, today)
        for row, e in zip(lo[1:], exp):
            got = dict(zip(lo[0], row))
            assert got["Days open"] == ("" if e["days"] == "" else "%g" % e["days"]), (got, e)
            assert got["Is open"] == e["open"] and got["Overdue"] == e["over"], (got, e)
        los = list(csv.reader(open(files["Summary"], encoding="utf-8")))
        cached = [[("" if _v(c) is None else _v(c)) for c in (r or [])] for r in _sheet(out, "Summary")["rows"]]
        n = 0
        for lr, cr in zip(los, cached):
            for j, c in enumerate(cr[1:4], start=1):
                if isinstance(c, (int, float)) and not isinstance(c, bool) and j < len(lr) and lr[0] != "Today":
                    assert abs(float(lr[j]) - float(c)) < 0.051, (lr, cr)
                    n += 1
                elif isinstance(c, str) and j < len(lr) and c and lr[0] not in ("Today",) and not c.startswith("See"):
                    assert lr[j] == c, (lr, cr)
        assert n >= 30, n
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ------------------------------------------------------------------ several files, one file, other kinds, options
def test_multi_file_order_and_single_file():
    a, b = _run(A), _run(B)
    ba = _tracker([b, a])
    heads, data = _data(ba)
    assert data[0]["Source file"].startswith(NBN)            # the first file's copy is kept
    # B first: A's copies hold Est hours, a column B does not have, so merging them into B's copy would lose a value;
    # they are kept, flagged as the same ID, never merged (nothing is dropped, nothing is filled in)
    mg = [r for r in _sheet(ba, "Merged duplicates")["rows"][1:] if r[0].startswith("across")]
    assert mg == [], mg
    t115 = [d for d in data if _v(d["ID"]) == "T-115"]
    assert len(t115) == 2 and all("a column only one file holds" in d["Check"] for d in t115), t115
    assert ba["summary"]["rows"] == 43
    one = _tracker([a])
    assert one["summary"]["rows"] == 26 and one["summary"]["merged_across"] == 0
    assert [f["name"] for f in one["summary"]["files"]] == [NA]
    only_b = json.loads(NB.tracker_json(json.dumps({"sha256": b})))
    assert only_b["ok"] and [f["name"] for f in only_b["summary"]["files"]] == [NBN]
    assert {x["sha256"] for x in only_b["available"]} >= {a, b}
    gone = json.loads(NB.tracker_json(json.dumps({"sha256": "0" * 64})))
    assert not gone["ok"] and "run it again" in gone["error"]
    # Track E's several-file intake: deliver_json('tracker', settings) calls build(input_frames(), settings); a frame
    # names its file by the summary's sha256, and the tracker reads that file's cleaner view (not its kept rows)
    frames = [{"id": "in1", "name": NA, "primary": True, "role": "tasks", "df": None, "kinds": {},
               "summary": {"sha256": a, "ok": True, "name": NA}},
              {"id": "in2", "name": NBN, "primary": False, "role": "tasks", "df": None, "kinds": {},
               "summary": {"sha256": b, "ok": True, "name": NBN}},
              {"id": "in3", "name": "broken.csv", "primary": False, "role": "other", "summary": {"ok": False, "error": "empty"}}]
    ef = TR.build(frames, {"kind": "", "questions": "", "dup_key": "exact"})
    assert ef["ok"] and ef["summary"]["rows"] == 37 and ef["summary"]["merged_across"] == 6 and ef["today"] == AS_OF, ef.get("error")
    un = _sheet(ef, "Unplaced")["rows"][1:]
    assert any(r[0] == "broken.csv" and r[1] is None and "empty" in r[3] for r in un), un
    assert TR.build(frames, {"dup_key": "id"})["summary"]["merged_across"] == 9


def test_duplicate_key_merges_and_says_what_differed():
    out = _pair(dup_key=["ID"])
    mg = [r for r in _sheet(out, "Merged duplicates")["rows"][1:] if r[0].startswith("across")]
    assert len(mg) == 9 and all(r[0] == "across files: the same ID" for r in mg)
    d117 = next(r for r in mg if "T-117" in r[4])
    assert d117[5] == "Status, Closed", d117
    assert out["summary"]["rows"] == 34


def test_kind_suggestion_and_templates():
    out = _pair()
    assert out["kind"] == "tasks" and "task" in out["suggest"]["why"]
    inv = b"SKU,Item,Category,Location,Qty on hand,Unit cost,Reorder level\nA-1,Bolts M6,Fixings,Bin 3,120,0.15,50\n" \
          b"A-2,Nuts M6,Fixings,Bin 3,40,0.08,50\nB-7,Gloves,PPE,Shelf 1,12,3.5,10\nB-8,Goggles,PPE,Shelf 1,,7.25,5\n"
    o = _tracker([_run("", inv, "stock-list.csv")])
    assert o["kind"] == "inventory", o["suggest"]
    h, d = _data(o)
    assert "Stock value" in h and "Reorder" in h
    assert [_v(r["Stock value"]) for r in d] == [18.0, 3.2, 42.0, ""]
    assert [_v(r["Reorder"]) for r in d] == ["No", "Yes", "No", ""]
    sv = next(r for r in _sheet(o, "Summary")["rows"] if r and _v(r[0]) == "Stock value (Qty x Unit cost, summed)")
    assert abs(_v(sv[1]) - 63.2) < 1e-9 and sv[1]["f"].startswith("SUM(Data!")
    shifts = b"Date,Staff,Role,Start,End,Status\n2026-03-02,Team A,Front desk,09:00,17:00,worked\n" \
             b"2026-03-02,Team B,Cleaner,22:00,06:00,worked\n2026-03-03,Team A,Front desk,09:00,,scheduled\n"
    o = _tracker([_run("", shifts, "rota-export.csv")])
    assert o["kind"] == "shifts", o["suggest"]
    h, d = _data(o)
    assert [_v(r["Hours"]) for r in d] == [8.0, 8.0, ""], [_v(r["Hours"]) for r in d]
    bills = b"Supplier,Invoice no,Category,Invoice date,Amount,Paid?\nAcme Supplies,INV-1,Parts,2026-01-05,120.50,yes\n" \
            b"Acme Supplies,INV-2,Parts,2026-02-05,80,no\nBright Clean,BC-9,Cleaning,2026-02-10,300,no\n"
    o = _tracker([_run("", bills, "bills.csv")])
    assert o["kind"] == "vendors", o["suggest"]
    h, d = _data(o)
    assert [_v(r["Vendor total"]) for r in d] == [200.5, 200.5, 300.0]
    up = next(r for r in _sheet(o, "Summary")["rows"] if r and str(_v(r[0])).startswith("Amount not yet paid"))
    assert _v(up[1]) == 380.0, up
    # the visitor chooses another kind: same rows, that kind's columns
    o2 = _tracker([_run(A)], kind="requests")
    assert o2["kind"] == "requests" and _data(o2)[0][:3] == ["ID", "Request", "Status"]


def test_unknown_status_is_kept_flagged_and_listed():
    data = b"Task,Status,Opened\nFix sink,done,2026-03-01\nPaint wall,awaiting parts,2026-03-02\nMop,open,2026-03-03\n"
    o = _tracker([_run("", data, "todo-list.csv")])
    h, d = _data(o)
    assert [_v(r["Status"]) for r in d][1].lower() == "awaiting parts"
    assert "not one of the tracker's status words" in d[1]["Check"]
    lists = _sheet(o, "Lists")["rows"]
    assert any(_v(r[0]).lower() == "awaiting parts" for r in lists[1:] if r), lists


def test_cleaning_edits_flow_into_the_tracker():
    a = _run(A)
    o = _tracker([a], rules={"rules": [{"name": "type_case", "mode": "title"}]})
    h, d = _data(o)
    assert {_v(r["Category"]) for r in d} == {"Maintenance", "Safety", "Admin", "Finance"}


TESTS = [v for k, v in sorted(globals().items()) if k.startswith("test_")]

if __name__ == "__main__":
    bad = 0
    for t in TESTS:
        try:
            t()
            print("PASS", t.__name__)
        except Exception as e:  # noqa: BLE001
            bad += 1
            import traceback
            traceback.print_exc()
            print("FAIL", t.__name__, e)
    print("%d of %d passed" % (len(TESTS) - bad, len(TESTS)))
    sys.exit(1 if bad else 0)
