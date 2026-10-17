#!/usr/bin/env python3
"""Tests for the monthly reporting pack (Track E; engine/nl_pack.py, src/js/61-pack.js, CONTRACT §5.18) and the
several-files intake every step-6 report type reads (nl_browser.input_json / input_frames / deliver_json).

    NL_ENGINE_ROOT=<copy of northledger-core> python tools/test_nl_pack.py

Checks, on the SYNTHETIC pack demo (tools/make_sample_pack.py: the sales ledger + its budget + a KPI extract) and on
small files built here:
  * every figure of the pack (this month, last month, change, plan, variance, year to date) equals a recomputation
    with pandas from the raw CSV files (a different code path: read_csv + groupby, the engine's set-aside rows dropped
    by their source line), and the planted facts are found
  * levels, counts at a date and rates are never summed over months (no YTD; a stock at the month's latest date);
    a rate's change is in percentage points; gaps stay gaps
  * the watch list is exactly the lines over the threshold (percent and amount bases, several thresholds), and where a
    flagged move sits matches pandas (the variance by region; the month's change by region through the engine's
    change step)
  * the Excel workbook made by the page's own code (node tools/pack_files.mjs: 56-xlsx.js + 61-pack.js) holds visible
    formulas in Metrics and Watch list; evaluated here by a small formula evaluator, every formula gives the engine's
    figure, and its cached value is that figure; the watch flags follow ONE threshold cell (changing it changes them)
  * the Word cover note: one page (LibreOffice to PDF, when soffice is installed), the reader named, the watch list,
    "driver to be confirmed"; no detailed variance commentary
  * the recipe round trip: the recipe's settings make the same pack; its file name patterns match next month's names;
    a column the recipe read that is gone is reported
  * several files: each companion read by the same run, the main run's state (Ask your data's frame, Clean my file's
    record) unchanged after it, the same file twice read once, a new main file clears them, roles guessed; a scenario
    column (Actual/Budget), paired columns, a long table read in its own shape, a trial balance; the reporting month
    steps back when the month is not complete
Prints PASS/FAIL per test, exits 1 on any failure.
"""
from __future__ import annotations

import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import warnings
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
SITE = os.path.normpath(os.path.join(HERE, ".."))
ENGINE_ROOT = os.path.abspath(os.environ.get("NL_ENGINE_ROOT") or os.path.join(SITE, "..", "northledger-core"))
sys.path.insert(0, ENGINE_ROOT)
sys.path.insert(0, os.path.join(SITE, "engine"))
os.environ.setdefault("NL_BROWSER_STRICT", "1")
warnings.filterwarnings("ignore")

import pandas as pd  # noqa: E402

import nl_browser as nb  # noqa: E402
import nl_pack as npk  # noqa: E402

AS_OF = "2026-10-01"
E = os.path.join(SITE, "engine")
LEDGER, BUDGET, KPI = "sample-sales-synthetic.csv", "sample-pack-budget-synthetic.csv", "sample-pack-kpi-synthetic.csv"
TMP = tempfile.mkdtemp(prefix="nl_pack_test_")
_DEMO = {}


def read(name):
    with open(os.path.join(E, name), "rb") as f:
        return f.read()


def main_file(data, name, as_of=AS_OF):
    """A new main file, as the worker's scan starts one: the last run's other files go first."""
    _DEMO.clear()
    nb.inputs_reset()
    return json.loads(nb.run_json(data, name, "", None, as_of))


def demo(settings=None, fresh=False):
    """The demo run: the ledger as the main file, the budget and the KPI extract as companions; then the pack."""
    if fresh or "rep" not in _DEMO:
        rep = main_file(read(LEDGER), LEDGER)
        _DEMO["rep"] = rep
        _DEMO["in2"] = json.loads(nb.input_json(read(BUDGET), BUDGET, AS_OF))
        _DEMO["in3"] = json.loads(nb.input_json(read(KPI), KPI, AS_OF))
    s = {"as_of": AS_OF}
    s.update(settings or {})
    return json.loads(nb.deliver_json("monthly-pack", json.dumps(s)))


def line(pack, key):
    return next(x for x in pack["lines"] if x["key"] == key)


def close(a, b, tol=1e-6):
    if a is None or b is None:
        return a is None and b is None
    return abs(a - b) <= tol * max(1.0, abs(a), abs(b))


# ----------------------------------------------------------------------------- pandas, independently
def ledger_kept():
    """The raw ledger without the rows the engine set aside (by their source line: the line number in the file)."""
    rep = _DEMO["rep"]
    q = pd.read_csv(io.StringIO(rep["downloads"]["quarantine_csv"]))
    drop = set(int(x) for x in q["source_line"]) if len(q) else set()
    d = pd.read_csv(os.path.join(E, LEDGER))
    d["source_line"] = range(2, len(d) + 2)
    d = d[~d["source_line"].isin(drop)].copy()
    d["m"] = d["order_date"].str[:7]
    return d


def budget():
    b = pd.read_csv(os.path.join(E, BUDGET))
    b = b[b["month"] <= AS_OF[:7]]                 # the engine sets aside months after the analysis date
    return b


def kpi():
    return pd.read_csv(os.path.join(E, KPI))


# ----------------------------------------------------------------------------- a small formula evaluator
class Book:
    """Evaluates the workbook's formulas (the subset the pack writes: + - * /, comparisons, IF, OR, AND, ABS, ISNUMBER,
    cell and cross-sheet references) from the values openpyxl reads, so a formula is checked by what it computes."""

    def __init__(self, path):
        import openpyxl
        self.wb = openpyxl.load_workbook(path, data_only=False)
        self.cached = openpyxl.load_workbook(path, data_only=True)
        self.memo = {}
        self.override = {}

    def cell(self, sheet, ref):
        ref = ref.replace("$", "")
        key = (sheet, ref)
        if key in self.override:
            return self.override[key]
        if key in self.memo:
            return self.memo[key]
        v = self.wb[sheet][ref].value
        if isinstance(v, str) and v.startswith("="):
            v = self.eval(sheet, v[1:])
        elif v is None:
            v = ""
        self.memo[key] = v
        return v

    def eval(self, sheet, f):
        toks = re.findall(r"'(?:[^']|'')+'!\$?[A-Z]+\$?\d+|[A-Z]+\$?\d+|\$[A-Z]+\$?\d+|\"[^\"]*\"|\d+\.?\d*(?:[eE][-+]?\d+)?|[A-Z]+(?=\()|<>|<=|>=|[-+*/(),=<>]", f)
        keep = getattr(self, "t", None), getattr(self, "i", 0), getattr(self, "sheet", None)   # a formula of a formula
        self.t, self.i, self.sheet = toks, 0, sheet
        try:
            v = self.comparison()
            assert self.i == len(self.t), (f, self.t[self.i:])
        finally:
            self.t, self.i, self.sheet = keep
        return v

    def peek(self):
        return self.t[self.i] if self.i < len(self.t) else None

    def take(self):
        self.i += 1
        return self.t[self.i - 1]

    def comparison(self):
        a = self.additive()
        while self.peek() in ("=", "<>", "<", ">", "<=", ">="):
            op = self.take()
            b = self.additive()
            if op == "=":
                a = (a == b) if not (a == "" and b == 0) else True
            elif op == "<>":
                a = a != b
            else:
                a = {"<": a < b, ">": a > b, "<=": a <= b, ">=": a >= b}[op]
        return a

    @staticmethod
    def n(v):
        return 0 if v == "" else v

    def additive(self):
        a = self.term()
        while self.peek() in ("+", "-"):
            op = self.take()
            b = self.term()
            a = self.n(a) + self.n(b) if op == "+" else self.n(a) - self.n(b)
        return a

    def term(self):
        a = self.unary()
        while self.peek() in ("*", "/"):
            op = self.take()
            b = self.unary()
            if op == "/" and self.n(b) == 0:
                a = "#DIV/0!"                          # Excel's error value (an IF branch not taken may hold one)
                continue
            a = "#DIV/0!" if "#DIV/0!" in (a, b) else (self.n(a) * self.n(b) if op == "*" else self.n(a) / self.n(b))
        return a

    def unary(self):
        if self.peek() == "-":
            self.take()
            return -self.n(self.unary())
        return self.atom()

    def atom(self):
        t = self.take()
        if t == "(":
            v = self.comparison()
            assert self.take() == ")"
            return v
        if t.startswith('"'):
            return t[1:-1]
        if re.match(r"^\d", t):
            return float(t)
        if "!" in t:
            sh, ref = t.rsplit("!", 1)
            return self.cell(sh.strip("'").replace("''", "'"), ref)
        if self.peek() == "(" and t in ("IF", "OR", "AND", "ABS", "ISNUMBER"):
            self.take()
            args = []
            if t == "IF":
                c = self.comparison()
                self.take()
                a = self.comparison()
                self.take()
                b = self.comparison()
                assert self.take() == ")"
                return a if c else b
            while True:
                args.append(self.comparison())
                s = self.take()
                if s == ")":
                    break
            if t == "OR":
                return any(args)
            if t == "AND":
                return all(args)
            if t == "ABS":
                return abs(self.n(args[0]))
            return isinstance(args[0], (int, float)) and not isinstance(args[0], bool)
        if re.match(r"^\$?[A-Z]+\$?\d+$", t):
            return self.cell(self.sheet, t)
        raise AssertionError("unknown token %r" % t)


def pack_files(pack, made="2026-10-01"):
    d = os.path.join(TMP, "files-%d" % len(os.listdir(TMP)))
    pj = os.path.join(TMP, "pack-%d.json" % len(os.listdir(TMP)))
    with open(pj, "w") as f:
        json.dump(pack, f)
    r = subprocess.run(["node", os.path.join(SITE, "tools", "pack_files.mjs"), pj, d, made], capture_output=True, text=True, timeout=120)
    assert r.returncode == 0, r.stderr[-2000:]
    return d, json.loads(r.stdout)


# ----------------------------------------------------------------------------- tests
def test_figures_match_pandas():
    pack = demo()
    assert pack["ok"], pack["error"]
    assert pack["month"] == "2026-09" and pack["prev_month"] == "2026-08", pack["month"]
    d = ledger_kept()
    b = budget()
    bb = b[b["scenario"] == "Budget"]
    for col in ("revenue", "cost", "units"):
        L = line(pack, col)
        cur = d.loc[d["m"] == "2026-09", col].sum()
        prv = d.loc[d["m"] == "2026-08", col].sum()
        plan = bb.loc[bb["month"] == "2026-09", col].sum()
        assert close(L["cur"], cur) and close(L["prev"], prv), (col, L["cur"], cur, L["prev"], prv)
        assert close(L["plan"], plan), (col, L["plan"], plan)
        assert close(L["mom"], cur - prv) and close(L["mom_pct"], (cur - prv) / abs(prv), 1e-5)
        assert close(L["var"], cur - plan) and close(L["var_pct"], (cur - plan) / abs(plan), 1e-5)
        ytd = [("2026-%02d" % m) for m in range(1, 10)]
        ya, yp = d.loc[d["m"].isin(ytd), col].sum(), bb.loc[bb["month"].isin(ytd), col].sum()
        assert close(L["ytd_act"], ya, 1e-6) and close(L["ytd_plan"], yp), (col, L["ytd_act"], ya, L["ytd_plan"], yp)
        assert close(L["ytd_var"], ya - yp, 1e-5)
    # gross margin from the month's sums, never an average of rates; its change in points
    g = line(pack, "gross_margin")
    s9, s8 = d[d["m"] == "2026-09"], d[d["m"] == "2026-08"]
    m9 = (s9.revenue.sum() - s9.cost.sum()) / s9.revenue.sum()
    m8 = (s8.revenue.sum() - s8.cost.sum()) / s8.revenue.sum()
    assert close(g["cur"], m9, 1e-5) and close(g["prev"], m8, 1e-5) and close(g["mom"], (m9 - m8) * 100, 1e-4), g
    assert g["kind"] == "rate" and g["scale"] == 100.0 and g["derived"]["op"] == "ratio"
    k = kpi().set_index("month")
    for col in ("active_customers", "new_customers", "churn_rate_pct", "nps", "headcount", "support_tickets"):
        L = line(pack, col)
        assert close(L["cur"], float(k.loc["2026-09", col])) and close(L["prev"], float(k.loc["2026-08", col])), (col, L)
        tgt = col + "_target"
        if tgt in k.columns:
            assert close(L["plan"], float(k.loc["2026-09", tgt])), (col, L["plan"])
        else:
            assert L["plan"] is None and L["planned"] is None and L["var"] is None


def test_planted_facts():
    pack = demo()
    rev, churn, nps, tick, head = (line(pack, k) for k in ("revenue", "churn_rate_pct", "nps", "support_tickets", "headcount"))
    assert rev["var_pct"] < -0.10 and rev["watch"], rev["var_pct"]
    assert rev["driver"] and rev["driver"]["by"] == "region" and rev["driver"]["member"] == "EMEA", rev["driver"]
    assert abs(rev["driver"]["amount"]) > abs(rev["var"]), "EMEA carries more than the whole shortfall"
    assert close(churn["mom"], 1.2, 1e-6) and close(churn["var"], 1.4, 1e-6) and churn["watch"], churn
    assert close(nps["mom"], -7.0) and nps["points"] and nps["watch"]
    assert close(tick["mom_pct"], (545 - 414) / 414, 1e-5) and tick["watch"] and tick["plan"] is None
    assert head["mom"] == 0 and not head["watch"] and head["kind"] == "stock"
    assert "2026-07" in nps["gaps"], nps["gaps"]
    assert any("NPS has no value for July 2026" in w for w in pack["notes"]["wrong"]), pack["notes"]["wrong"]
    for L in pack["lines"]:
        assert L["cause"] == npk.CAUSE


def test_levels_never_summed():
    pack = demo()
    for k in ("active_customers", "headcount", "nps", "churn_rate_pct", "gross_margin"):
        L = line(pack, k)
        assert L["kind"] != "flow", (k, L["kind"])
        if k != "gross_margin":
            assert L["ytd_act"] is None and L["ytd_plan"] is None and L["ytd_why"], (k, L)
    assert line(pack, "new_customers")["ytd_act"] == float(kpi().set_index("month").loc[[("2026-%02d" % m) for m in range(1, 10)], "new_customers"].sum())
    # a stock with several rows a month is taken at the month's latest date, added across those rows, never summed over dates
    rows = ["date,region,headcount,orders"]
    for m in range(1, 10):
        for day in (5, 28):
            for reg, h in (("North", 10), ("South", 20)):
                rows.append("2026-%02d-%02d,%s,%d,%d" % (m, day, reg, h + m + (day == 28) * 2, 3 + m))
    # (not strict: the analyst loop's change step raises on a count at a date split by a group with no entity column,
    # nl_analyst.op_change; the report stands without that block, as it does in the browser)
    os.environ.pop("NL_BROWSER_STRICT", None)
    try:
        main_file(("\n".join(rows) + "\n").encode(), "staff.csv")
    finally:
        os.environ["NL_BROWSER_STRICT"] = "1"
    p = json.loads(nb.deliver_json("monthly-pack", json.dumps({"as_of": AS_OF})))
    assert p["ok"], p["error"]
    hc, od = line(p, "headcount"), line(p, "orders")
    # September: the 28th's rows (the month's latest date), North 10+9+2 and South 20+9+2; August likewise with 8
    assert hc["kind"] == "stock" and hc["cur"] == 21 + 31 and hc["prev"] == 20 + 30 and hc["ytd_act"] is None, hc
    assert od["kind"] == "flow" and od["cur"] == 4 * 12 and od["prev"] == 4 * 11 and od["ytd_act"] == sum(4 * (3 + m) for m in range(1, 10)), od


def test_watch_list_against_threshold():
    for thr in (5, 10, 20, 50):
        pack = demo({"threshold": thr})
        want = []
        for L in pack["lines"]:
            if any(v is not None and abs(v) > thr / 100.0 for v in (L["mom_pct"], L["var_pct"])):
                want.append(L["key"])
        assert pack["watch"] == want, (thr, pack["watch"], want)
    pack = demo({"threshold": 20000, "basis": "amount"})
    want = [L["key"] for L in pack["lines"] if any(v is not None and abs(v) > 20000 for v in (L["mom"], L["var"]))]
    assert pack["watch"] == want and "revenue" in want, (pack["watch"], want)
    # where the month's change sits (the engine's change-by-group step), checked with pandas: at 5% revenue's -6.6% month
    # is flagged on the change
    pack = demo({"threshold": 5})
    rev = line(pack, "revenue")
    assert rev["watch_why"][0] == "the month's change", rev["watch_why"]
    d = rev["driver"]
    assert d and d["of"] == "change", d
    k = ledger_kept()
    g = k[k["m"].isin(["2026-08", "2026-09"])].pivot_table(index=d["by"], columns="m", values="revenue", aggfunc="sum")
    delta = g["2026-09"] - g["2026-08"]
    assert close(d["amount"], float(delta[d["member"]]), 1e-6) and close(d["total"], float(delta.sum()), 1e-6), (d, delta)
    # the variance by region: EMEA's actual minus EMEA's budget
    pack = demo()
    d = line(pack, "revenue")["driver"]
    b = budget()
    va = k[k["m"] == "2026-09"].groupby("region")["revenue"].sum() - b[(b["month"] == "2026-09") & (b["scenario"] == "Budget")].groupby("region")["revenue"].sum()
    assert d["of"] == "variance" and close(d["amount"], float(va["EMEA"]), 1e-6) and close(d["total"], float(va.sum()), 1e-6), (d, va)


def test_forecast_member_and_month_choice():
    pack = demo({"compare": "forecast"})
    b = budget()
    fc = b[(b["scenario"] == "Forecast") & (b["month"] == "2026-09")]
    assert close(line(pack, "revenue")["plan"], fc["revenue"].sum()), line(pack, "revenue")["plan"]
    # no forecast for January: the year to date plan is not made up
    assert line(pack, "revenue")["ytd_plan"] is None and "no plan figure" in line(pack, "revenue")["ytd_why"]
    pack = demo({"month": "2026-08"})
    assert pack["month"] == "2026-08" and pack["month_why"] == "you chose it"


def test_workbook_formulas():
    pack = demo()
    d, files = pack_files(pack)
    names = [f["name"] for f in files]
    assert names == ["monthly-pack-2026-09.xlsx", "monthly-pack-2026-09-cover-note.docx", "monthly-pack-2026-09-recipe.json",
                     "monthly-pack-2026-09-checklist.md"], names
    bk = Book(os.path.join(d, names[0]))
    assert bk.wb.sheetnames == ["Metrics", "Watch list", "Notes", "Inputs checklist"], bk.wb.sheetnames
    M = bk.wb["Metrics"]
    head = [c.value for c in M[5]]
    assert head[:6] == ["Line", "Kind", "Last month (August 2026)", "This month (September 2026)", "Change", "Change %"], head
    cols = {"E": "mom", "F": "mom_pct", "H": "var", "I": "var_pct", "L": "ytd_var", "M": "ytd_var_pct"}
    nform = 0
    for i, L in enumerate(pack["lines"]):
        r = 6 + i
        assert M["A%d" % r].value == L["label"]
        for c, k in cols.items():
            f = M["%s%d" % (c, r)].value
            assert isinstance(f, str) and f.startswith("="), (L["key"], c, f)   # a visible formula
            nform += 1
            got = bk.cell("Metrics", "%s%d" % (c, r))
            want = L[k]
            num = isinstance(got, (int, float)) and not isinstance(got, bool)
            assert (got == "" and want is None) or (num and want is not None and close(got, want, 1e-6)), (L["key"], c, f, got, want)
            cached = bk.cached["Metrics"]["%s%d" % (c, r)].value
            assert (cached in (None, "") and want is None) or close(float(cached), want, 1e-9), (L["key"], c, cached, want)
        if L.get("derived"):
            assert str(M["D%d" % r].value).startswith("="), "a derived line's values are formulas of the revenue and cost rows"
            assert close(bk.cell("Metrics", "D%d" % r), L["cur"], 1e-6)
    W = bk.wb["Watch list"]
    assert W["A2"].value == "Threshold" and close(W["B2"].value, 0.10)
    flags = {}
    for i, L in enumerate(pack["lines"]):
        r = 6 + i
        f = W["D%d" % r].value
        assert f.startswith("=IF(OR(") and "$B$2" in f, f
        flags[L["key"]] = bk.cell("Watch list", "D%d" % r)
        assert (flags[L["key"]] == "WATCH") == L["watch"], (L["key"], flags[L["key"]])
        assert (bk.cached["Watch list"]["D%d" % r].value or "") == ("WATCH" if L["watch"] else "")
        assert (W["G%d" % r].value == npk.CAUSE) == L["watch"]
    # ONE threshold cell: 30% in B2 and the flags follow it
    bk.memo.clear()
    bk.override[("Watch list", "B2")] = 0.30
    p30 = demo({"threshold": 30})
    for i, L in enumerate(pack["lines"]):
        assert (bk.cell("Watch list", "D%d" % (6 + i)) == "WATCH") == (L["key"] in p30["watch"]), L["key"]
    assert nform >= 6 * len(pack["lines"])
    N = [r[0].value for r in bk.wb["Notes"].iter_rows()]
    assert "Looked wrong, please check" in N and "Assumptions" in N and "What each line is" in N
    C = bk.wb["Inputs checklist"]
    assert [C.cell(row=5 + i, column=2).value for i in range(3)] == [LEDGER, KPI, BUDGET]


def test_cover_note():
    pack = demo({"readers": "the board"})
    d, files = pack_files(pack)
    docx = os.path.join(d, files[1]["name"])
    with zipfile.ZipFile(docx) as z:
        x = z.read("word/document.xml").decode("utf-8")
    text = re.sub(r"<[^>]+>", " ", x)
    for w in ("Monthly pack: September 2026", "for the board", "Headline movements", "Watch list", "driver to be confirmed",
              "SYNTHETIC", "EMEA"):
        assert w in text, w
    spec = json.load(open(os.path.join(d, "specs.json")))["cover"]
    kinds = [b["type"] for b in spec["blocks"]]
    assert kinds.count("h2") == 2 and "pagebreak" not in kinds and spec["style"] == "business", kinds
    soffice = shutil.which("soffice")
    if not soffice:
        print("  (soffice not installed: the one-page check did not run)")
        return
    prof = os.path.join(TMP, "lo-profile")
    r = subprocess.run([soffice, "-env:UserInstallation=file://" + prof, "--headless", "--convert-to", "pdf", "--outdir", d, docx],
                       capture_output=True, text=True, timeout=180)
    pdf = docx[:-5] + ".pdf"
    assert os.path.exists(pdf), r.stderr[-500:]
    pages = len(re.findall(rb"/Type\s*/Page[^s]", open(pdf, "rb").read()))
    assert pages == 1, "the cover note is %d pages" % pages


def test_recipe_round_trip():
    pack = demo({"readers": "the CFO", "threshold": 12, "metrics": "revenue, churn, nps"})
    assert [L["key"] for L in pack["lines"]] == ["revenue", "churn_rate_pct", "nps"], [L["key"] for L in pack["lines"]]
    d, files = pack_files(pack)
    rec = json.load(open(os.path.join(d, files[2]["name"])))
    assert rec["kind"] == "monthly-pack" and rec["settings"]["threshold"] == 12 and rec["settings"]["readers"] == "the CFO"
    again = demo(dict(rec["settings"], recipe=rec))
    for k in ("month", "watch", "headline", "threshold", "basis", "compare", "readers"):
        assert again[k] == pack[k], (k, again[k], pack[k])
    assert [L["key"] for L in again["lines"]] == [L["key"] for L in pack["lines"]]
    assert {x["pattern"] for x in rec["inputs"]} == {LEDGER, BUDGET, KPI}
    assert npk._match("kpi-*-*.csv", "kpi-2026-10.csv") and not npk._match("kpi-*-*.csv", "budget-2026-10.csv")
    assert npk.name_pattern("kpi-extract-2026-09.csv") == "kpi-extract-*-*.csv"
    # next month's KPI extract lost a column the recipe read: said, not guessed
    k = kpi().drop(columns=["nps_target"])
    main_file(read(LEDGER), LEDGER)
    nb.input_json(read(BUDGET), BUDGET, AS_OF)
    nb.input_json(k.to_csv(index=False).encode(), KPI, AS_OF)
    p2 = json.loads(nb.deliver_json("monthly-pack", json.dumps({"as_of": AS_OF, "recipe": rec})))
    assert any("nps_target" in w and "recipe" in w for w in p2["notes"]["wrong"]), p2["notes"]["wrong"]
    md = open(os.path.join(d, files[3]["name"])).read()
    assert "October 2026" in md and "`order_date`" in md and "`scenario`" in md, md[:400]
    demo(fresh=True)


def test_several_files():
    rep = main_file(read(LEDGER), LEDGER)
    assert [x["id"] for x in rep["inputs"]] == ["in1"] and rep["inputs"][0]["primary"] and rep["inputs"][0]["role"] == "actuals"
    frame_cols = list(nb._LAST_ANALYST["df"].columns)
    clean_sha = nb._LAST_CLEAN.get("sha256")
    a = json.loads(nb.input_json(read(BUDGET), BUDGET, AS_OF))
    b = json.loads(nb.input_json(read(KPI), KPI, AS_OF))
    assert list(nb._LAST_ANALYST["df"].columns) == frame_cols and nb._LAST_CLEAN.get("sha256") == clean_sha, \
        "a companion changed the main run's state"
    assert a["role"] == "plan" and a["scenario"]["column"] == "scenario" and \
        {m["value"]: m["reads_as"] for m in a["scenario"]["members"]} == {"Budget": "plan", "Forecast": "forecast"}, a
    assert a["rows_set_aside"] == 12 and a["rows_kept"] == 42, (a["rows_kept"], a["rows_set_aside"])
    assert b["role"] == "actuals_vs_plan" and {p["stem"] for p in b["pairs"]} == {"active_customers", "new_customers", "churn_rate_pct", "nps", "headcount"}
    dup = json.loads(nb.input_json(read(KPI), "copy.csv", AS_OF))
    assert dup.get("duplicate") and len(json.loads(nb.inputs_json())) == 3
    assert [x["id"] for x in nb.input_frames()] == ["in1", "in2", "in3"]
    roles = {x["id"]: x["role"] for x in nb.input_frames({"in3": "kpi"})}
    assert roles["in3"] == "kpi"
    nb.inputs_reset()
    assert len(json.loads(nb.inputs_json())) == 1
    # a scenario column with Actual and Budget in one long file, read in its own shape (the run reads it one series at a time)
    rows = ["month,line,scenario,amount"]
    for m in range(1, 10):
        for ln, base in (("Revenue", 1000), ("Payroll", 400)):
            act = base * (1.25 if (m == 9 and ln == "Payroll") else 1.0) + m
            rows += ["2026-%02d,%s,Actual,%d" % (m, ln, act), "2026-%02d,%s,Budget,%d" % (m, ln, base + m)]
    rep = main_file(("\n".join(rows) + "\n").encode(), "pl-2026-09.csv")
    s = rep["inputs"][0]
    assert s["role"] == "actuals_vs_plan" and s["rows_kept"] == 36 and s.get("own_shape"), s
    p = json.loads(nb.deliver_json("monthly-pack", json.dumps({"as_of": AS_OF})))
    pay = line(p, "payroll")
    assert pay["cur"] == 509 and pay["plan"] == 409 and pay["watch"] and close(pay["var_pct"], 100 / 409, 1e-5), pay
    # a trial balance: debit minus credit per account
    tb = "month,account,debit,credit\n2026-08,Sales,0,5000\n2026-08,Rent,800,0\n2026-09,Sales,0,5600\n2026-09,Rent,800,0\n"
    t = json.loads(nb.input_json(tb.encode(), "tb-2026-09.csv", AS_OF))
    assert t["role"] == "trial_balance", t["role"]
    p = json.loads(nb.deliver_json("monthly-pack", json.dumps({"as_of": AS_OF})))
    sl = line(p, "sales")
    assert sl["cur"] == -5600 and sl["prev"] == -5000 and line(p, "rent")["mom"] == 0
    assert not any(L["key"] == "gross_margin" for L in p["lines"]), "a margin is never made of two files' lines"
    # a file of daily rows that stops mid-month: that month is not complete, the pack reports the month before
    d = pd.read_csv(os.path.join(E, LEDGER))
    d = d[d["order_date"] <= "2026-09-16"]
    main_file(d.to_csv(index=False).encode(), "ledger.csv", "2026-09-16")
    p = json.loads(nb.deliver_json("monthly-pack", json.dumps({"as_of": "2026-09-16"})))
    assert p["month"] == "2026-08" and "not complete" in p["month_why"], (p["month"], p["month_why"])
    # no monthly figures at all: said plainly
    main_file(b"name,score\nA,1\nB,2\nC,3\n", "list.csv")
    p = json.loads(nb.deliver_json("monthly-pack", "{}"))
    assert not p["ok"] and "date column" in p["error"], p
    assert json.loads(nb.deliver_json("nonsense", "{}"))["ok"] is False
    demo(fresh=True)


def test_record_is_json_and_figure_free_words():
    pack = demo()
    json.dumps(pack, allow_nan=False)
    for L in pack["lines"]:
        assert L["label"] and not re.search(r"\d{3,}", L["label"]), L["label"]
    assert pack["recipe"]["settings"]["metrics"] == []


TESTS = [v for k, v in sorted(globals().items()) if k.startswith("test_")]


def main() -> int:
    bad = 0
    try:
        for t in sorted(TESTS, key=lambda f: f.__code__.co_firstlineno):
            try:
                t()
                print("PASS", t.__name__)
            except Exception as e:  # noqa: BLE001
                bad += 1
                import traceback
                print("FAIL", t.__name__, "%s: %s" % (type(e).__name__, str(e)[:600]))
                traceback.print_exc(limit=2)
    finally:
        shutil.rmtree(TMP, ignore_errors=True)
    print("%d passed, %d failed" % (len(TESTS) - bad, bad))
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
